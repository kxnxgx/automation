# -*- coding: utf-8 -*-
"""
Excelシート生成モジュール (generate_sheets.py)
============================================
"""

import os
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from config import AutomationError, BrandConfig
from utils import (
    find_single_csv, is_target_brand, parse_csv_store_structure,
    get_brand_name_safe
)
from data_loader import get_store_sales_value

# ============================================================
# 手順1: 出荷予定振分.csv から RETAIL ワークブックの作成
# ============================================================
def step1_create_retail_wb(input_dir, encoding, brand_config: BrandConfig):
    print("[手順1] 出荷予定振分.csv から RETAILシートの構築")
    src = find_single_csv(input_dir, ["出荷予定振分"])
    df = pd.read_csv(src, encoding=encoding, header=0, dtype=str)
    
    # ======== フォーマット変動検証は廃止済み ========
    # 店舗数は parse_csv_store_structure() で動的に取得するため、
    # 固定列数を前提とする検証ブロックは不要。
    # ================================================
    
    # ======== 自ブランドのみ抽出フィルタリング ========
    if len(df) > 3:
        df_headers = df.iloc[:3].copy()
        df_data = df.iloc[3:].copy()
        
        # 4列目(iloc[3])が商品コード、6列目(iloc[5])がBrand
        mask = df_data.apply(lambda row: is_target_brand(brand_config, row.iloc[3], row.iloc[5]), axis=1)
        df_data = df_data[mask].copy()
        
        df = pd.concat([df_headers, df_data], ignore_index=True)
    # ================================================
    
    wb_retail = openpyxl.Workbook()
    # 【修正④】デフォルトフォントを9ptに設定（完全版に合わせる）
    # ワークブック作成直後にfont index 0を変更することで、
    # 明示的なフォント指定がない全セルが9ptになる（全セル走査不要）
    wb_retail._fonts[0] = openpyxl.styles.Font(name='Calibri', size=9)
    ws = wb_retail.active
    ws.title = "RETAIL"
    
    # ヘッダー書き込み
    for c_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c_idx).value = col_name
        
    # データ書き込み
    for r_idx, row in enumerate(df.values.tolist(), start=2):
        for c_idx, val in enumerate(row, start=1):
            if pd.isna(val) if isinstance(val, float) else val is None:
                ws.cell(row=r_idx, column=c_idx).value = None
            else:
                if isinstance(val, str):
                    try:
                        if val.isdigit():
                            ws.cell(row=r_idx, column=c_idx).value = int(val)
                        else:
                            ws.cell(row=r_idx, column=c_idx).value = float(val)
                    except ValueError:
                        ws.cell(row=r_idx, column=c_idx).value = val
                else:
                    ws.cell(row=r_idx, column=c_idx).value = val
                    
    return wb_retail

# ============================================================
# 手順2: Python直接集計・マージおよび書式適用
# ============================================================
def step2_python_direct_merge(wb_retail, base_dir, input_dir, template_path, brand_config: BrandConfig,
                              pivot_df, zozo_dict, oioi_dict, ec_dict, hutte_dict,
                              zozo_m_dict, oioi_m_dict, tokka_dict):
    print("[手順2] Pythonによる直接マージ＆書式適用 (VBA代替処理)")
    
    if "order" in wb_retail.sheetnames:
        del wb_retail["order"]

    ws_retail = wb_retail["RETAIL"]
    
    # 店舗構造を解析して、RETAILシートでの出荷指示エリアの開始列を動的に決定する
    csv_struct = parse_csv_store_structure(input_dir, "cp932")
    N_struct = csv_struct['N']
    gap1_struct = csv_struct['gap1']
    gap2_struct = csv_struct['gap2']
    
    # 出荷指示エリアの開始列 = 2 * N + gap1 + gap2 + 9 (通常は49列目/AW列)
    # CSV構造: [0-5]=固定6列, [6]=NWA, [7]=納品先BULK, [8..8+N-1]=BULK店舗,
    #   gap1列, 在庫N列, gap2列, 出荷指示数N列 → 出荷指示数開始インデックス = 8+N+gap1+N+gap2
    #   Excel列番号 = インデックス+1 = 2*N + gap1 + gap2 + 9
    retail_col_start = 2 * N_struct + gap1_struct + gap2_struct + 9
    
    # 本来のCSV由来の列数をヘッダー行（1行目）から動的に特定する
    csv_cols_count = 0
    for c in range(1, ws_retail.max_column + 1):
        if ws_retail.cell(row=1, column=c).value is not None:
            csv_cols_count = c
    if csv_cols_count == 0:
        csv_cols_count = 69 # フォールバック
    
    # 追加データ列以降をクリアする。
    # 行1 (CSV列番号行) 以降 retail_col_start 以降はすべてクリアして再設定する。
    # CSV由来のヘッダー (出荷指示数・出荷予定日) や列番号が残留するのを防ぐ。
    max_r_retail = max(ws_retail.max_row, 1)
    clear_end_col = max(retail_col_start + 40, csv_cols_count + 1)
    
    # 【修正③】クリア前に1行目（CSV列番号インデックス）の値を保存しておく
    # retail_col_start 以降の列インデックス値は完全版に合わせて保持・復元する
    saved_row1 = {}
    for c in range(retail_col_start, clear_end_col):
        saved_row1[c] = ws_retail.cell(row=1, column=c).value
    
    for r in range(1, max_r_retail + 1):  # 行1 (CSV列番号行) 以降すべてクリア
        for c in range(retail_col_start, clear_end_col):
            ws_retail.cell(row=r, column=c).value = None
    
    # 【修正③】クリア後にrow1のインデックス値を復元する（完全版との一致）
    for c, val in saved_row1.items():
        ws_retail.cell(row=1, column=c).value = val
            
    # シートの複製
    ws_order = wb_retail.copy_worksheet(ws_retail)
    ws_order.title = "order"

    # "商品コード" "商品名" "ブランド名" の列を動的に特定
    idx_code = None
    idx_name = None
    idx_brand = None
    
    for c in range(1, ws_order.max_column + 1):
        val = ws_order.cell(row=1, column=c).value
        val_str = str(val).strip() if val is not None else ""
        if val_str in ("商品コード", "4"):
            idx_code = c
        elif val_str in ("商品名", "5"):
            idx_name = c
        elif val_str in ("6", "ブランド名", "表記部門名1"):
            idx_brand = c
            
    if not idx_code: idx_code = 4
    if not idx_name: idx_name = 5
    if not idx_brand: idx_brand = 6

    # 動的な列削除
    cols_to_delete_before = idx_code - 1
    if cols_to_delete_before > 0:
        ws_order.delete_cols(1, cols_to_delete_before)
        idx_name -= cols_to_delete_before
        idx_brand -= cols_to_delete_before
        idx_code = 1
        
    # A列の書式適用
    for row in range(1, ws_order.max_row + 1):
        cell = ws_order.cell(row=row, column=1)
        cell.number_format = 'General'
        cell.alignment = Alignment(horizontal='left', vertical='center')

    # B列幅設定
    ws_order.column_dimensions['B'].width = 15
    
    # ブランド名/部門名が入っているC列を削除
    if idx_brand > 0:
        ws_order.delete_cols(idx_brand, 1)

    ws_order["D4"].value = "RETAIL確保"
    ws_order.delete_rows(3, 1) # 3行目を削除 (店舗コード行を削除)
    
    # ============================================================
    # 店舗構造の動的取得
    # ============================================================
    csv_struct = parse_csv_store_structure(input_dir, "cp932")
    N    = csv_struct['N']
    gap1 = csv_struct['gap1']
    gap2 = csv_struct['gap2']
    csv_store_names = csv_struct['store_names']
    csv_store_codes = csv_struct['store_codes']
    
    # RETAILシートの出荷予定日の列 = 出荷指示数開始列 + N店舗分 (動的に計算)
    retail_date_col = retail_col_start + N
    
    # 1. BULK不要列の削除 (gap1列分。BULK店舗終端の次から)
    ws_order.delete_cols(N + 5, gap1)
    # 2. 在庫不要列の削除 (gap2列分。在庫店舗終端の次から)
    ws_order.delete_cols(2 * N + 5, gap2)
    
    # 列定数導出
    kakubo_start_col   = 3
    kakubo_end_col     = N + 4
    zaiko_start_col    = N + 5
    zaiko_end_col      = 2 * N + 4
    order_start_col    = 2 * N + 5
    order_end_col      = 3 * N + 4
    kabusoku_start_col = 3 * N + 5
    kabusoku_end_col   = 4 * N + 4
    
    kakubo_start_letter   = openpyxl.utils.get_column_letter(kakubo_start_col)
    kakubo_end_letter     = openpyxl.utils.get_column_letter(kakubo_end_col)
    zaiko_start_letter    = openpyxl.utils.get_column_letter(zaiko_start_col)
    zaiko_end_letter      = openpyxl.utils.get_column_letter(zaiko_end_col)
    order_start_letter    = openpyxl.utils.get_column_letter(order_start_col)
    order_end_letter      = openpyxl.utils.get_column_letter(order_end_col)
    kabusoku_start_letter = openpyxl.utils.get_column_letter(kabusoku_start_col)
    kabusoku_end_letter   = openpyxl.utils.get_column_letter(kabusoku_end_col)

    # "確保" ヘッダー
    ws_order.merge_cells(f"{kakubo_start_letter}2:{kakubo_end_letter}2")
    ws_order[f"{kakubo_start_letter}2"].value = "確保"
    ws_order[f"{kakubo_start_letter}2"].alignment = Alignment(horizontal='center', vertical='center')
    fill_kakubo = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_blue_c2_d3 = PatternFill(start_color="00CCFFFF", end_color="00CCFFFF", fill_type="solid") # 水色
    for r in range(2, 4):
        for c in range(kakubo_start_col, kakubo_end_col + 1):
            ws_order.cell(row=r, column=c).fill = fill_kakubo
    for r in range(2, 4):
        for c in range(kakubo_start_col, kakubo_start_col + 2):
            ws_order.cell(row=r, column=c).fill = fill_blue_c2_d3

    # "在庫" ヘッダー
    ws_order.merge_cells(f"{zaiko_start_letter}2:{zaiko_end_letter}2")
    ws_order[f"{zaiko_start_letter}2"].value = "在庫"
    ws_order[f"{zaiko_start_letter}2"].alignment = Alignment(horizontal='center', vertical='center')
    fill_zaiko = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    for r in range(2, 4):
        for c in range(zaiko_start_col, zaiko_end_col + 1):
            ws_order.cell(row=r, column=c).fill = fill_zaiko

    # "ORDER" ヘッダー
    ws_order.merge_cells(f"{order_start_letter}2:{order_end_letter}2")
    ws_order[f"{order_start_letter}2"].value = "ORDER"
    ws_order[f"{order_start_letter}2"].alignment = Alignment(horizontal='center', vertical='center')
    fill_order = PatternFill(start_color="FFFF00", fill_type="solid")
    for r in range(2, 4):
        for c in range(order_start_col, order_end_col + 1):
            ws_order.cell(row=r, column=c).fill = fill_order

    # 過不足エリアのクリア
    max_r_order = max(ws_order.max_row, 1)
    for r in range(1, max_r_order + 1):
        for c in range(kabusoku_start_col, kabusoku_end_col + 1):
            ws_order.cell(row=r, column=c).value = None

    # "過不足" ヘッダー
    ws_order.merge_cells(f"{kabusoku_start_letter}2:{kabusoku_end_letter}2")
    ws_order[f"{kabusoku_start_letter}2"].value = "過不足"
    ws_order[f"{kabusoku_start_letter}2"].alignment = Alignment(horizontal='center', vertical='center')

    # 3行目の見出しを縦書きに
    align_vertical = Alignment(textRotation=255, horizontal='center', vertical='center')
    for c in range(kakubo_start_col, order_end_col + 1):
        ws_order.cell(row=3, column=c).alignment = align_vertical
        col_letter = openpyxl.utils.get_column_letter(c)
        ws_order.column_dimensions[col_letter].width = 5

    # A2:B2 切り貼り
    ws_order["A3"].value = ws_order["A2"].value
    ws_order["B3"].value = ws_order["B2"].value
    ws_order["A2"].value = None
    ws_order["B2"].value = None
    ws_order["A3"].alignment = Alignment(horizontal='center', vertical='center')
    ws_order["B3"].alignment = Alignment(horizontal='center', vertical='center')

    # 最終行検出
    lastRow = 4
    for r in range(ws_order.max_row, 3, -1):
        if ws_order.cell(row=r, column=2).value is not None:
            lastRow = r
            break

    # BULKの確保店舗列を非表示
    for c in range(5, kakubo_end_col + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws_order.column_dimensions[col_letter].hidden = True

    # --------------------------------------------------------
    # 店舗名を dynamic store names から各セクション (確保, 在庫, ORDER, 過不足) に適用
    # --------------------------------------------------------
    print("  店舗名を各エリア（確保、在庫、ORDER、過不足）の3行目に設定中...")
    for c_idx in range(N):
        store_name = csv_store_names[c_idx]
        # 確保 (BULK) column starts at 5
        ws_order.cell(row=3, column=5 + c_idx).value = store_name
        # 在庫 starts at zaiko_start_col
        ws_order.cell(row=3, column=zaiko_start_col + c_idx).value = store_name
        # ORDER starts at order_start_col
        ws_order.cell(row=3, column=order_start_col + c_idx).value = store_name
        # 過不足 starts at kabusoku_start_col
        ws_order.cell(row=3, column=kabusoku_start_col + c_idx).value = store_name

    # --------------------------------------------------------
    # 列インデックスを動的に検索・解決する仕組みの構築
    # --------------------------------------------------------
    col_map = {}
    current_area = None
    for col in range(1, ws_order.max_column + 1):
        area_val = ws_order.cell(row=2, column=col).value
        if area_val is not None:
            area_val_str = str(area_val).strip()
            if area_val_str in ("確保", "在庫", "ORDER", "過不足"):
                current_area = area_val_str
        
        store_val = ws_order.cell(row=3, column=col).value
        if store_val is not None and current_area is not None:
            store_val_str = str(store_val).strip()
            col_map[(current_area, store_val_str)] = col

    stores_list = csv_store_names  # CSVから動的に取得した店舗名リスト
    
    bulk_store_map = {}
    order_store_map = {}
    kabusoku_store_map = {}
    
    for (area, store), col_idx in col_map.items():
        matched_store = None
        for s in stores_list:
            norm_s = "OIOI" if s == "丸井web" else ("EC" if s in ("EC", "YSEC") else s)
            norm_store = "OIOI" if store == "丸井web" else ("EC" if store in ("EC", "YSEC") else store)
            if norm_s == norm_store:
                matched_store = s
                break
        
        if matched_store:
            std_key = "OIOI" if matched_store == "丸井web" else ("EC" if matched_store in ("EC", "YSEC") else matched_store)
            if area == "確保":
                bulk_store_map[std_key] = col_idx
            elif area == "ORDER":
                order_store_map[std_key] = col_idx
            elif area == "過不足":
                kabusoku_store_map[std_key] = col_idx

    col_c = 3
    col_d = 4
    col_kabusoku_total = None
    col_retail_kakubo = None
    col_kabusoku_zan = None
    
    for c in range(1, ws_order.max_column + 1):
        val = ws_order.cell(row=3, column=c).value
        val_str = str(val).strip() if val is not None else ""
        if c > order_end_col:
            if "合計" in val_str or val_str == "過不足合計":
                col_kabusoku_total = c
            elif val_str == "RETAIL確保" and c != col_d:
                col_retail_kakubo = c
            elif (val_str == "過不足" or val_str == "過不足残") and c not in kabusoku_store_map.values():
                col_kabusoku_zan = c

    if not col_kabusoku_total: col_kabusoku_total = kabusoku_start_col + N
    if not col_retail_kakubo: col_retail_kakubo  = kabusoku_start_col + N + 1
    if not col_kabusoku_zan:  col_kabusoku_zan   = kabusoku_start_col + N + 2

    # --------------------------------------------------------
    # 商品マスタ (商品マスタ全項目.csv) のロード
    # --------------------------------------------------------
    print("  商品マスタ全項目.csv からコードと商品名を取得中...")
    master_dict = {}
    master_path = os.path.join(base_dir, "商品マスタ全項目.csv")
    if os.path.exists(master_path):
        try:
            df_m = pd.read_csv(master_path, encoding="cp932", dtype=str)
            code_col = None
            name_col = None
            for col in df_m.columns:
                col_str = str(col).strip()
                if "商品コード" in col_str or "品番" in col_str:
                    code_col = col
                elif "商品名" in col_str or "品名" in col_str:
                    name_col = col
            if code_col is None: code_col = df_m.columns[0]
            if name_col is None: name_col = df_m.columns[2] if len(df_m.columns) > 2 else df_m.columns[1]

            def _normalize_product_code(v):
                """指数表記のコード (例: 1.22002E+12) を整数文字列に正規化する。"""
                if v is None or (isinstance(v, float) and pd.isna(v)):
                    return ""
                s = str(v).strip()
                # 指数表記を検出して整数化
                if 'E' in s.upper() or 'e' in s:
                    try:
                        return str(int(float(s)))
                    except (ValueError, OverflowError):
                        pass
                return s

            for _, row in df_m.iterrows():
                m_code = _normalize_product_code(row[code_col] if not pd.isna(row[code_col]) else None)
                m_name = str(row[name_col]).strip() if not pd.isna(row[name_col]) else ""
                if m_code:
                    master_dict[m_code] = {
                        'name': m_name,
                        'code': m_code
                    }
        except Exception as e:
            print(f"  警告: 商品マスタのロードに失敗しました: {e}")

    # --------------------------------------------------------
    # 売上CSVからバックアップ用の「コード ↔ 商品名」を取得
    # --------------------------------------------------------
    backup_names = {}
    try:
        src_store = find_single_csv(input_dir, ["営業日付別売上分析"])
        df_store_raw = pd.read_csv(src_store, encoding="cp932", dtype=str)
        for _, row in df_store_raw.iterrows():
            code = str(row["3rd Item No."]).strip() if not pd.isna(row["3rd Item No."]) else ""
            name = str(row["商品名"]).strip() if "商品名" in row and not pd.isna(row["商品名"]) else ""
            if code and name:
                backup_names[code] = name
    except Exception: pass

    try:
        src_zozo = find_single_csv(input_dir, ["卸売上明細"], exclude_keywords=["卸売上明細 (1)", "卸売上明細(1)"])
        df_zozo_raw = pd.read_csv(src_zozo, encoding="cp932", dtype=str)
        for _, row in df_zozo_raw.iterrows():
            code = str(row["商品コード"]).strip() if not pd.isna(row["商品コード"]) else ""
            name = str(row["商品名"]).strip() if "商品名" in row and not pd.isna(row["商品名"]) else ""
            if code and name:
                backup_names[code] = name
    except Exception: pass

    try:
        src_oioi = find_single_csv(input_dir, ["卸売上明細 (1)"])
        df_oioi_raw = pd.read_csv(src_oioi, encoding="cp932", dtype=str)
        for _, row in df_oioi_raw.iterrows():
            code = str(row["商品コード"]).strip() if not pd.isna(row["商品コード"]) else ""
            name = str(row["商品名"]).strip() if "商品名" in row and not pd.isna(row["商品名"]) else ""
            if code and name:
                backup_names[code] = name
    except Exception: pass

    try:
        src_ec = find_single_csv(input_dir, ["order_"])
        df_ec_raw = pd.read_csv(src_ec, encoding="cp932", dtype=str)
        key_col = "オプション独自コード" if "オプション独自コード" in df_ec_raw.columns else "商品コード"
        for _, row in df_ec_raw.iterrows():
            code = str(row[key_col]).strip() if not pd.isna(row[key_col]) else ""
            name = str(row["商品名"]).strip() if "商品名" in row and not pd.isna(row["商品名"]) else ""
            if code and name:
                backup_names[code] = name
    except Exception: pass

    # --------------------------------------------------------
    # 売上があるすべての商品コードの抽出
    # --------------------------------------------------------
    all_sales_codes = set()
    for code, qty in zozo_dict.items():
        if qty != 0: all_sales_codes.add(code)
    for code, qty in oioi_dict.items():
        if qty != 0: all_sales_codes.add(code)
    for code, qty in ec_dict.items():
        if qty != 0 and "GIFT" not in code:
            all_sales_codes.add(code)
    if pivot_df is not None and not pivot_df.empty:
        num_cols = [c for c in pivot_df.columns if c != "3rd Item No."]
        for _, row in pivot_df.iterrows():
            code = str(row["3rd Item No."]).strip() if not pd.isna(row["3rd Item No."]) else ""
            if code:
                total_qty = pd.to_numeric(row[num_cols], errors="coerce").fillna(0).sum()
                if total_qty != 0:
                    all_sales_codes.add(code)
                    
    # HUTTE出荷指示に存在する指示コードも追加 (ZOZO/OIOIマスタは属性情報のため行追加の対象外)
    for code, qty in hutte_dict.items():
        if qty != 0: all_sales_codes.add(code)


    # 出荷予定振分の現在のA列コード
    order_codes = set()
    for r in range(4, lastRow + 1):
        val = ws_order.cell(row=r, column=1).value
        if val:
            order_codes.add(str(val).strip())

    leaked_codes = sorted(list(all_sales_codes - order_codes))
    print(f"  出荷予定にない売上発生商品を追加中... (追加数: {len(leaked_codes)})")

    # --------------------------------------------------------
    # orderシートおよびRETAILシートの末尾への行追加
    # --------------------------------------------------------
    lastRowRetail = 4
    for r in range(ws_retail.max_row, 3, -1):
        if ws_retail.cell(row=r, column=4).value is not None:
            lastRowRetail = r
            break

    current_order_row = lastRow + 1
    current_retail_row = lastRowRetail + 1

    start_order_col = min(order_store_map.values()) if order_store_map else 31
    end_order_col = col_kabusoku_zan

    for code in leaked_codes:
        # 商品コードがこのブランドに属さない場合はスキップ
        # （ブランド名列を持たない売上CSVによる他ブランド商品の混入を防ぐ防衛的チェック）
        if not is_target_brand(brand_config, code, None):
            continue
        p_name = "不明な商品"
        if code in master_dict and master_dict[code]['name']:
            p_name = master_dict[code]['name']
        elif code in backup_names:
            p_name = backup_names[code]
            
        p_code = code
        if code in master_dict and master_dict[code]['code']:
            p_code = master_dict[code]['code']

        ws_order.cell(row=current_order_row, column=1).value = p_code
        ws_order.cell(row=current_order_row, column=2).value = p_name
        
        for col in range(kakubo_start_col, kakubo_end_col + 1):
            ws_order.cell(row=current_order_row, column=col).value = 0
        for col in range(zaiko_start_col, zaiko_end_col + 1):
            ws_order.cell(row=current_order_row, column=col).value = 0
            
        for store_name, col_idx in order_store_map.items():
            if store_name == "ZOZO":
                val = int(zozo_dict.get(code, 0))
            elif store_name == "OIOI":
                val = int(oioi_dict.get(code, 0))
            elif store_name in ("YSEC", "EC"):
                val = int(ec_dict.get(code, 0))
            else:
                val = get_store_sales_value(pivot_df, code, store_name)
            ws_order.cell(row=current_order_row, column=col_idx).value = max(0, val)

        ws_retail.cell(row=current_retail_row, column=4).value = p_code
        ws_retail.cell(row=current_retail_row, column=5).value = p_name
        ws_retail.cell(row=current_retail_row, column=6).value = brand_config.name
        ws_retail.cell(row=current_retail_row, column=7).value = 0
        for col in range(8, 29):
            ws_retail.cell(row=current_retail_row, column=col).value = 0
            
        for c_idx in range(N):
            target_col_letter = openpyxl.utils.get_column_letter(retail_col_start + c_idx)
            source_col_letter = openpyxl.utils.get_column_letter(order_start_col + c_idx)
            ws_retail.cell(row=current_retail_row, column=retail_col_start + c_idx).value = f"=order!{source_col_letter}{current_order_row}"
            
        ws_retail.cell(row=current_retail_row, column=retail_date_col).value = int(datetime.datetime.now().strftime("%Y%m%d"))

        current_order_row += 1
        current_retail_row += 1

    lastRow = current_order_row - 1
    lastRowRetail = current_retail_row - 1

    start_zaiko_col = min(col_idx for (area, store), col_idx in col_map.items() if area == "在庫")
    end_zaiko_col = max(col_idx for (area, store), col_idx in col_map.items() if area == "在庫")
    
    thin_side = Side(border_style="thin", color="000000")
    border_left_only = Border(left=thin_side)
    border_right_only = Border(right=thin_side)
    for r in range(2, lastRow + 1):
        ws_order.cell(row=r, column=start_zaiko_col).border = border_left_only
        ws_order.cell(row=r, column=end_zaiko_col).border = border_right_only

    # --------------------------------------------------------
    # 既存商品の売上データ書き込み
    # --------------------------------------------------------
    print("  既存商品の売上データを書き込み中...")
    for r in range(4, lastRow + 1):
        code_short = ws_order.cell(row=r, column=1).value
        if not code_short:
            for store_name, col_idx in order_store_map.items():
                ws_order.cell(row=r, column=col_idx).value = 0
            continue
            
        code_short = str(code_short).strip()

        raw_sales = {}
        for store_name, col_idx in order_store_map.items():
            qty = 0
            if store_name == "ZOZO":
                qty = int(zozo_dict.get(code_short, 0))
            elif store_name == "OIOI":
                qty = int(oioi_dict.get(code_short, 0))
            elif store_name in ("YSEC", "EC"):
                qty = int(ec_dict.get(code_short, 0))
            else:
                qty = get_store_sales_value(pivot_df, code_short, store_name)
            raw_sales[store_name] = max(0, qty)

        total_sales_all = sum(raw_sales.values())

        def _safe_num(v):
            if isinstance(v, (int, float)): return v
            try: return float(v)
            except (ValueError, TypeError): return 0

        c_val = _safe_num(ws_order.cell(row=r, column=3).value)
        d_val = _safe_num(ws_order.cell(row=r, column=4).value)
        kabu_sum = sum(_safe_num(ws_order.cell(row=r, column=col).value) for col in range(5, kakubo_end_col + 1))

        if c_val == 0 and d_val == 0 and kabu_sum < total_sales_all:
            for store_name, col_idx in order_store_map.items():
                ws_order.cell(row=r, column=col_idx).value = 0
        else:
            for store_name, col_idx in order_store_map.items():
                ws_order.cell(row=r, column=col_idx).value = raw_sales[store_name]

    # --------------------------------------------------------
    # 計算式（過不足、過不足残、SUBTOTAL合計）の書き込み
    # --------------------------------------------------------
    print("  過不足などの計算式を設定中...")
    for r in range(4, lastRow + 1):
        for s in stores_list:
            std_s = "OIOI" if s == "丸井web" else ("EC" if s in ("EC", "YSEC") else s)
            b_col = bulk_store_map.get(std_s)
            o_col = order_store_map.get(std_s)
            k_col = kabusoku_store_map.get(std_s)
            if b_col and o_col and k_col:
                b_let = openpyxl.utils.get_column_letter(b_col)
                o_let = openpyxl.utils.get_column_letter(o_col)
                ws_order.cell(row=r, column=k_col).value = f"={b_let}{r}-{o_let}{r}"
            
        kabusoku_start_let = openpyxl.utils.get_column_letter(min(kabusoku_store_map.values()))
        kabusoku_end_let = openpyxl.utils.get_column_letter(max(kabusoku_store_map.values()))
        
        let_c = openpyxl.utils.get_column_letter(col_c)
        let_d = openpyxl.utils.get_column_letter(col_d)
        let_kabusoku_total = openpyxl.utils.get_column_letter(col_kabusoku_total)
        let_retail_kakubo = openpyxl.utils.get_column_letter(col_retail_kakubo)
        let_kabusoku_zan = openpyxl.utils.get_column_letter(col_kabusoku_zan)
        
        ws_order.cell(row=r, column=col_kabusoku_total).value = f"=SUMIF({kabusoku_start_let}{r}:{kabusoku_end_let}{r},\"<0\",{kabusoku_start_let}{r}:{kabusoku_end_let}{r})"
        ws_order.cell(row=r, column=col_retail_kakubo).value = f"=IF({let_kabusoku_total}{r}<0,{let_d}{r}+{let_kabusoku_total}{r},{let_d}{r})"
        nwa_formula = f"IF({let_retail_kakubo}{r}<0,{let_c}{r}+{let_retail_kakubo}{r},{let_c}{r})"
        ws_order.cell(row=r, column=col_kabusoku_zan).value = f'=IF({nwa_formula}=0,"",{nwa_formula})'

    for c in range(1, ws_order.max_column + 1):
        ws_order.cell(row=1, column=c).value = None
        
    max_order_col = max(order_store_map.values()) if order_store_map else order_end_col
    # 【修正①】SUBTOTAL範囲を3500固定（完全版に合わせ、フィルター適用後も全行を正しく合計できるよう）
    SUBTOTAL_MAX_ROW = 3500
    for c in range(3, max_order_col + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws_order.cell(row=1, column=c).value = f"=SUBTOTAL(9,{col_letter}4:{col_letter}{SUBTOTAL_MAX_ROW})"

    kabusoku_cols_indices = list(kabusoku_store_map.values()) + [col_kabusoku_total, col_retail_kakubo, col_kabusoku_zan]
    kabusoku_min_col = min(kabusoku_cols_indices) if kabusoku_cols_indices else 44
    kabusoku_max_col = max(kabusoku_cols_indices) if kabusoku_cols_indices else 59
    
    for c in range(kabusoku_min_col, kabusoku_max_col + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws_order.column_dimensions[col_letter].width = 4.43
        if c != col_kabusoku_zan:
            ws_order.column_dimensions[col_letter].hidden = True
            
    zan_letter = openpyxl.utils.get_column_letter(col_kabusoku_zan)
    ws_order[f"{zan_letter}3"].value = "過不足"
    ws_order[f"{zan_letter}3"].alignment = Alignment(textRotation=255, horizontal='center', vertical='center')
    ws_order[f"{zan_letter}3"].fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

    # 【修正②】条件付き書式を完全版に合わせる
    # 適用範囲: A1:A1048576（全行）、数式: BG1<0（行固定参照なし）
    ws_order.conditional_formatting._cf_rules.clear()
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    rule = FormulaRule(formula=[f"{zan_letter}1<0"], stopIfTrue=False, fill=red_fill)
    ws_order.conditional_formatting.add("A1:A1048576", rule)

    lastRowRetail = 5
    for r in range(ws_retail.max_row, 4, -1):
        if ws_retail.cell(row=r, column=4).value is not None:
            lastRowRetail = r
            break
            
    for r in range(5, lastRowRetail + 1):
        for c_idx in range(N):
            target_col_letter = openpyxl.utils.get_column_letter(retail_col_start + c_idx)
            source_col_letter = openpyxl.utils.get_column_letter(order_start_col + c_idx)
            source_row = r - 1
            ws_retail[f"{target_col_letter}{r}"] = f"=order!{source_col_letter}{source_row}"
        
        ws_retail.cell(row=r, column=retail_date_col).value = int(datetime.datetime.now().strftime("%Y%m%d"))
        
    num_slots = retail_date_col - retail_col_start
    
    for c_idx in range(num_slots):
        ws_retail.cell(row=2, column=retail_col_start + c_idx).value = "出荷指示数"
    ws_retail.cell(row=2, column=retail_date_col).value = "出荷予定日"
    
    for c_idx in range(N):
        if c_idx < len(csv_store_names):
            store_name = csv_store_names[c_idx]
            store_code = csv_store_codes[c_idx]
            ws_retail.cell(row=4, column=retail_col_start + c_idx).value = store_name
            try:
                ws_retail.cell(row=3, column=retail_col_start + c_idx).value = int(store_code)
            except (ValueError, TypeError):
                ws_retail.cell(row=3, column=retail_col_start + c_idx).value = store_code

    extra_col_start = col_kabusoku_zan + 1
    col_kibou = extra_col_start
    col_hutte = extra_col_start + 1
    col_zozo  = extra_col_start + 2
    col_oioi  = extra_col_start + 3
    col_r     = extra_col_start + 4
    col_tokka = extra_col_start + 5

    filter_end_letter = openpyxl.utils.get_column_letter(col_tokka)
    ws_order.auto_filter.ref = f"A3:{filter_end_letter}{lastRow}"
    ws_order.freeze_panes = "A4"

    print("  新規列（希望品、HUTTE、ZOZO、OIOI、特価等）のデータを書き込み中...")
    ws_order.cell(row=3, column=col_kibou).value = "希望品"
    ws_order.cell(row=3, column=col_hutte).value = "HUTTE"
    ws_order.cell(row=3, column=col_zozo).value  = "ZOZO"
    ws_order.cell(row=3, column=col_oioi).value  = "OIOI"
    ws_order.cell(row=3, column=col_r).value     = "R"
    ws_order.cell(row=3, column=col_tokka).value = "特価"
    
    for c in range(col_kibou, col_tokka + 1):
        cell = ws_order.cell(row=3, column=c)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        col_letter = openpyxl.utils.get_column_letter(c)
        ws_order.column_dimensions[col_letter].width = 8

    final_max_row = current_order_row - 1
    max_order_col = max(order_store_map.values()) if order_store_map else order_end_col

    for r in range(4, final_max_row + 1):
        code_val = ws_order.cell(row=r, column=1).value
        code_str = str(code_val).strip() if code_val else ""
        
        ws_order.cell(row=r, column=col_hutte).value = hutte_dict.get(code_str, "")
        ws_order.cell(row=r, column=col_zozo).value = zozo_m_dict.get(code_str, "")
        ws_order.cell(row=r, column=col_oioi).value = oioi_m_dict.get(code_str, "")
        ws_order.cell(row=r, column=col_r).value = r - 3
        
        tokka_val = tokka_dict.get(code_str, "")
        if tokka_val == 0 or tokka_val == "0" or tokka_val == 0.0:
            tokka_val = ""
        ws_order.cell(row=r, column=col_tokka).value = tokka_val

        hide_zero_cols = [3, 4] + list(range(zaiko_start_col, max_order_col + 1))
        for c in hide_zero_cols:
            val = ws_order.cell(row=r, column=c).value
            if val == 0 or val == "0" or val == 0.0:
                ws_order.cell(row=r, column=c).value = None

    print("  マージおよび書式適用完了")
