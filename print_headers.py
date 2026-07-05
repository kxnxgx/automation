import pandas as pd
import openpyxl

# 出荷予定振分.csv の最初の2行を読み込む
df = pd.read_csv('input/出荷予定振分.csv', encoding='cp932', nrows=2)
with open('headers_out.txt', 'w', encoding='utf-8') as f:
    f.write("Columns of 出荷予定振分.csv:\n")
    f.write(", ".join(df.columns.tolist()) + "\n\n")
    
    # テンプレートExcelの計算式2026シートの3行目を読み込む
    wb_temp = openpyxl.load_workbook("★RETAIL_テンプレート.xlsx", data_only=True)
    ws_temp = wb_temp["計算式2026"]
    row3 = [str(ws_temp.cell(row=3, column=c).value) for c in range(1, ws_temp.max_column + 1)]
    f.write("Row 3 of ★RETAIL_テンプレート.xlsx (計算式2026):\n")
    f.write(", ".join(row3) + "\n")
