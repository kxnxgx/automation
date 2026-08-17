# -*- coding: utf-8 -*-
"""
出荷明細作成自動化 共有モジュール (automation_core.py)
===================================================
共通ロジック、バグ修正、およびリファクタリングを統合したコアモジュールです。
"""

import os
import sys
import traceback
import openpyxl
import pandas as pd
import glob

# モジュール群からのインポート (互換性維持および再エクスポート)
from config import BrandConfig, AutomationError, BRAND_FRV, BRAND_TEN, BRAND_HWG
from utils import show_error, show_info, parse_csv_store_structure, is_target_brand, get_brand_name_safe, is_store_match
from data_loader import (
    load_store_sales, load_zozo_sales, load_oioi_sales, load_ec_sales,
    load_hutte, load_zozo_master, load_tokka
)
from generate_sheets import step1_create_retail_wb, step2_python_direct_merge

# ============================================================
# メインパイプラインの実行
# ============================================================
def run_pipeline(brand_config: BrandConfig, output_file_name, input_dir, base_dir, template_path, error_log_path):
    print("=" * 50)
    print(f"出荷明細作成自動化処理 開始 [{brand_config.name}用] (リファクタリング版)")
    print("=" * 50)

    wb_retail = None
    success = False
    encoding = "cp932"

    try:
        # 各種CSVの直接読み込みと集計の実行
        pivot_df = load_store_sales(input_dir, encoding, brand_config)
        zozo_dict = load_zozo_sales(input_dir, encoding, brand_config)
        oioi_dict = load_oioi_sales(input_dir, encoding, brand_config)
        ec_dict = load_ec_sales(input_dir, encoding, brand_config)

        # 新規ファイルからのデータ抽出
        hutte_dict = load_hutte(input_dir, brand_config)
        zozo_m_dict, oioi_m_dict = load_zozo_master(input_dir, brand_config)
        tokka_dict = load_tokka(input_dir, brand_config)

        # 出荷予定振分からベースシートを作成
        wb_retail = step1_create_retail_wb(input_dir, encoding, brand_config)
        
        # 直接集計結果をマージして書式を設定
        step2_python_direct_merge(
            wb_retail, base_dir, input_dir, template_path, brand_config,
            pivot_df, zozo_dict, oioi_dict, ec_dict, hutte_dict,
            zozo_m_dict, oioi_m_dict, tokka_dict
        )
        
        # 最終保存
        output_dir = os.path.join(base_dir, "output")
        os.makedirs(output_dir, exist_ok=True)
        retail_xlsx_path = os.path.join(output_dir, output_file_name)
        
        if os.path.exists(retail_xlsx_path):
            try:
                os.remove(retail_xlsx_path)
            except PermissionError as pe:
                raise AutomationError(
                    "完成版Excelファイルが開いたままです",
                    f"【発生した状況】\n完成版Excel（{output_file_name}）がExcelなどのアプリケーションで開かれた状態になっています。\n\n"
                    f"【対処方法】\n開いている「{output_file_name}」のExcel画面を完全に閉じてから、もう一度実行してください。"
                ) from pe
            except Exception:
                pass
            
        print(f"\n{output_file_name} を書き出し中...")
        try:
            wb_retail.save(retail_xlsx_path)
        except PermissionError as pe:
            raise AutomationError(
                "完成版Excelファイルが開いたままです",
                f"【発生した状況】\n完成版Excel（{output_file_name}）がExcelなどのアプリケーションで開かれた状態になっています。\n\n"
                f"【対処方法】\n開いている「{output_file_name}」のExcel画面を完全に閉じてから、もう一度実行してください。"
            ) from pe
            
        wb_retail.close()
        wb_retail = None
        
        print(f"\n処理が正常に完了しました！\n-> {os.path.basename(retail_xlsx_path)}")
        success = True
        
    except AutomationError as ae:
        with open(error_log_path, "w", encoding="utf-8") as f:
            f.write(f"エラー分類: {ae.title}\n")
            f.write(f"エラー詳細:\n{ae.message}\n")
            f.write("=" * 60 + "\n")
            f.write(traceback.format_exc())
        print(f"\nエラー: {ae.title}\n{ae.message}\n詳細は error.log を確認してください。")
        show_error(ae.title, ae.message)
        
    except Exception as e:
        title = "プログラム実行中にエラーが発生しました"
        msg = f"【発生した状況】\nプログラムの実行中に予期しないエラーが発生しました。\n\nエラー内容: {e}\n\n【対処方法】\n同じフォルダにある「error.log」ファイルの内容をシステム管理者または開発者に提供し、調査を依頼してください。"
        with open(error_log_path, "w", encoding="utf-8") as f:
            f.write(f"エラー分類: {title}\n")
            f.write(f"エラー詳細:\n{msg}\n")
            f.write("=" * 60 + "\n")
            f.write(traceback.format_exc())
        print(f"\nエラー: {title}\n{msg}\n詳細は error.log を確認してください。")
        show_error(title, msg)
        
    finally:
        if wb_retail is not None:
            try: wb_retail.close()
            except Exception: pass

    if success:
        show_info(f"処理が正常に完了しました！\n\n{output_file_name} を作成しました。")
    return success

# ============================================================
# データ検証用 共通パイプライン
# ============================================================
def verify_pipeline(brand_config: BrandConfig, input_dir: str, result_path: str) -> bool:
    """元CSVと生成されたExcel(orderシート)の売上データを商品・チャネルレベルで網羅的に検証する"""
    import pandas as pd
    import openpyxl
    import openpyxl.utils
    import glob
    import os

    print("=" * 55)
    print(f"  検証対象ファイル: {os.path.basename(result_path)}")
    print(f"  対象ブランド: {brand_config.name}")
    print("=" * 55)

    if not os.path.exists(result_path):
        print(f"[エラー] 検証対象のExcelファイルが見つかりません: {result_path}")
        return False

    if not os.path.exists(input_dir):
        print(f"[エラー] 入力データディレクトリが見つかりません: {input_dir}")
        return False

    def find_csv_local(keywords, exclude=None):
        clean_keywords = [k.replace(" ", "").replace("　", "") for k in keywords]
        clean_exclude = [e.replace(" ", "").replace("　", "") for e in exclude] if exclude else []
        for p in glob.glob(os.path.join(input_dir, "*.csv")):
            name = os.path.basename(p)
            clean_name = name.replace(" ", "").replace("　", "")
            if all(k in clean_name for k in clean_keywords):
                if clean_exclude and any(e in clean_name for e in clean_exclude):
                    continue
                return p
        return None

    # CSVファイルの検索
    zozo_csv = find_csv_local(["卸売上明細"], exclude=["卸売上明細 (1)", "卸売上明細(1)"])
    oioi_csv = find_csv_local(["卸売上明細 (1)"])
    store_csv = find_csv_local(["営業日付別売上分析"])
    ec_csv = find_csv_local(["order_"])

    if not zozo_csv:
        print("[エラー] 卸売上明細 CSV が見つかりません。")
        return False
    if not store_csv:
        print("[エラー] 営業日付別売上分析 CSV が見つかりません。")
        return False
    if not ec_csv:
        print("[エラー] order_*.csv が見つかりません。")
        return False

    # 1. 元CSVデータのロードと集計
    # ZOZO
    df_zozo = pd.read_csv(zozo_csv, encoding="cp932", dtype=str)
    if not df_zozo.empty:
        mask = df_zozo.apply(lambda row: is_target_brand(brand_config, row["商品コード"], get_brand_name_safe(row)), axis=1)
        df_zozo = df_zozo[mask].copy()
    df_zozo["販売数量"] = pd.to_numeric(df_zozo["販売数量"], errors="coerce").fillna(0).clip(lower=0)
    zozo_sales = df_zozo.groupby("商品コード")["販売数量"].sum().to_dict()

    # OIOI
    oioi_sales = {}
    if oioi_csv:
        df_oioi = pd.read_csv(oioi_csv, encoding="cp932", dtype=str)
        def _get_brand_o(row): return row.get("Brand") or row.get("ブランド名") or None
        if not df_oioi.empty:
            if "商品コード" in df_oioi.columns:
                mask = df_oioi.apply(lambda row: is_target_brand(brand_config, row["商品コード"], _get_brand_o(row)), axis=1)
            else:
                mask = df_oioi.apply(lambda row: is_target_brand(brand_config, None, _get_brand_o(row)), axis=1)
            df_oioi = df_oioi[mask].copy()
        df_oioi["販売数量"] = pd.to_numeric(df_oioi["販売数量"], errors="coerce").fillna(0).clip(lower=0)
        if "商品コード" in df_oioi.columns:
            oioi_sales = df_oioi.groupby("商品コード")["販売数量"].sum().to_dict()

    # 店舗
    df_store = pd.read_csv(store_csv, encoding="cp932", dtype=str)
    if not df_store.empty:
        mask = df_store.apply(lambda row: is_target_brand(brand_config, row["3rd Item No."], row.get("表記部門名1")), axis=1)
        df_store = df_store[mask].copy()
    df_store["数量"] = pd.to_numeric(df_store["数量"], errors="coerce").fillna(0).clip(lower=0)
    store_pivot = df_store.pivot_table(index="3rd Item No.", columns="店舗名称", values="数量", aggfunc="sum").fillna(0)

    # EC
    df_ec = pd.read_csv(ec_csv, encoding="cp932", dtype=str)
    # CRASH-02対応: 列名が変更された CSV でも KeyError で落ちないよう存在チェックを追加
    if "決済方法(ステータス)" in df_ec.columns:
        df_ec = df_ec[~df_ec["決済方法(ステータス)"].str.contains("キャンセル", na=False)]
    if "注文者" in df_ec.columns:
        df_ec = df_ec[~df_ec["注文者"].str.contains("店舗客注", na=False)]
    df_ec["個数"] = pd.to_numeric(df_ec["個数"], errors="coerce").fillna(0).clip(lower=0)
    key_col = "オプション独自コード" if "オプション独自コード" in df_ec.columns else "商品コード"
    # BUG-06対応: GIFTコードを data_loader と同様に除外（大文字統一）
    df_ec = df_ec[~df_ec[key_col].fillna("").str.upper().str.contains("GIFT", na=False)]

    brand_col = "ブランド名" if "ブランド名" in df_ec.columns else None
    if not df_ec.empty:
        if brand_col:
            mask = df_ec.apply(lambda row: is_target_brand(brand_config, row[key_col], row[brand_col]), axis=1)
        else:
            mask = df_ec[key_col].apply(lambda c: is_target_brand(brand_config, c, None))
        df_ec = df_ec[mask]
    ec_sales = df_ec.groupby(key_col)["個数"].sum().to_dict()

    # Excelワークブックの読み込み
    wb = openpyxl.load_workbook(result_path, data_only=True)
    if "order" not in wb.sheetnames:
        print(f"[エラー] Excelファイルに 'order' シートが存在しません: {result_path}")
        wb.close()
        return False
    ws = wb["order"]

    # 2. 前提データのエラーチェック (事前に実行)
    formula_errors = []
    for r in range(4, ws.max_row + 1):
        code_val = ws.cell(row=r, column=1).value
        if not code_val:
            continue
        for col in range(3, 18):
            val = ws.cell(row=r, column=col).value
            if val is not None:
                val_str = str(val).strip()
                if val_str.startswith("#"):
                    col_letter = openpyxl.utils.get_column_letter(col)
                    col_header = ws.cell(row=3, column=col).value or ""
                    formula_errors.append(
                        f"行 {r}, 列 {col_letter} ({col_header}): エラー値 '{val_str}'"
                    )

    if formula_errors:
        print("\n[WARNING/ERROR] 前提データに数式エラーが検出されました:")
        for err in formula_errors:
            print(f"  - {err}")
        wb.close()
        return False

    # 3. Excel列インデックスの動的解決 (列ズレ脆弱性の廃止)
    csv_struct = parse_csv_store_structure(input_dir, "cp932")
    N = csv_struct["N"]
    kakubo_end_col = N + 4
    all_names = csv_struct["store_names"]
    exclude_channels = {"ZOZO", "OIOI", "丸井web", "EC", "YSEC"}
    stores_list = [name for name in all_names if name not in exclude_channels]

    channels = ["ZOZO", "OIOI", "EC"] + stores_list
    col_map = {}

    current_area = None
    for col in range(1, ws.max_column + 1):
        area_val = ws.cell(row=2, column=col).value
        if area_val is not None:
            area_val_str = str(area_val).strip()
            if area_val_str in ("確保", "在庫", "ORDER", "過不足"):
                current_area = area_val_str

        if current_area != "ORDER":
            continue

        val = ws.cell(row=3, column=col).value
        if not val:
            continue
        val_str = str(val).strip()

        # チャネルごとに判定
        if "ZOZO" in val_str:
            col_map["ZOZO"] = col
        elif "OIOI" in val_str or "丸井" in val_str:
            col_map["OIOI"] = col
        elif "EC" in val_str or "YSEC" in val_str:
            col_map["EC"] = col
        else:
            for s in stores_list:
                if is_store_match(val_str, s):
                    col_map[s] = col
                    break

    missing_channels = [c for c in channels if c not in col_map]
    if missing_channels:
        print(f"[エラー] 以下のチャネルがExcelのヘッダー行から特定できませんでした: {missing_channels}")
        wb.close()
        return False

    # 4. 売上ゼロ化ロジックの適用と正解値（期待値）の算出
    resolved_sales = {c: {} for c in channels}

    def _safe_num_v(v):
        """セル値を安全に数値変換する（generate_sheets._safe_num と同一ロジック）。
        ⚠ BUG-03対応: data_only=True で開いたブックでは数式セルが None になる場合がある。
          また文字列型の数式（"=SUM(...)"）が残存する場合も 0 として扱う。
          在庫チェック列（C・D・BULK列）は CSV 由来の数値なので通常 None にはならないが、
          フォーマット変更等への防衛的対策として適用する。
        """
        if isinstance(v, (int, float)):
            return v
        if v is None:
            return 0
        try:
            return float(v)
        except (ValueError, TypeError):
            return 0

    for r in range(4, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if not code:
            continue
        code = str(code).strip()

        raw_vals = {}
        raw_vals["ZOZO"] = int(zozo_sales.get(code, 0))
        raw_vals["OIOI"] = int(oioi_sales.get(code, 0))
        raw_vals["EC"] = int(ec_sales.get(code, 0))

        # 店舗
        for s in stores_list:
            cols_match = [c for c in store_pivot.columns if is_store_match(c, s)]

            qty = 0
            if cols_match and code in store_pivot.index:
                qty = int(store_pivot.loc[code, cols_match].sum())
            raw_vals[s] = max(0, qty)

        # 全チャネルの売上合計で在庫不足を判定（仕様修正：全チャネルをゼロ化）
        total_sales_all = sum(max(0, raw_vals[k]) for k in channels)

        # 在庫情報（generate_sheets._safe_num と同一の変換ロジックを適用）
        c_val = _safe_num_v(ws.cell(row=r, column=3).value)
        d_val = _safe_num_v(ws.cell(row=r, column=4).value)
        kabu_sum = sum(_safe_num_v(ws.cell(row=r, column=col).value) for col in range(5, kakubo_end_col + 1))

        # 不足による売上ゼロ化（ZOZO/OIOI/EC/YSEC含む全チャネルをゼロにする）
        if c_val == 0 and d_val == 0 and kabu_sum < total_sales_all:
            for k in channels:
                resolved_sales[k][code] = 0
        else:
            for k in channels:
                resolved_sales[k][code] = max(0, raw_vals[k])


    # 正解合計値
    zozo_true = sum(resolved_sales["ZOZO"].values())
    oioi_true = sum(resolved_sales["OIOI"].values())
    ec_true = sum(resolved_sales["EC"].values())
    store_totals = {s: sum(resolved_sales[s].values()) for s in stores_list}

    # 5. Excel上の実際値の集計 (動的インデックス使用)
    def col_total(col_num):
        return sum(
            (ws.cell(row=r, column=col_num).value or 0)
            for r in range(4, ws.max_row + 1)
            if ws.cell(row=r, column=2).value is not None
        )

    actual = {c: col_total(col_map[c]) for c in channels}

    # 6. 商品レベル（行レベル）の検証
    diff_details = []
    for r in range(4, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if not code:
            continue
        code = str(code).strip()

        for c in channels:
            truth_val = resolved_sales[c].get(code, 0)
            col_idx = col_map[c]
            actual_val = ws.cell(row=r, column=col_idx).value or 0

            # 不一致があった場合
            if truth_val != actual_val:
                diff_details.append({
                    "row": r,
                    "code": code,
                    "channel": c,
                    "expected": truth_val,
                    "actual": actual_val
                })

    wb.close()

    # 7. 検証結果の出力
    print("\n" + "=" * 55)
    print("[1] 元CSV の合計（正解値）")
    print("=" * 55)
    print(f"  ZOZO    (卸売上明細.csv):          {zozo_true:>6}")
    print(f"  OIOI    (卸売上明細 (1).csv):      {oioi_true:>6}")
    for name, total in store_totals.items():
        print(f"  {name:<12}(営業日付別売上分析.csv): {total:>6}")
    print(f"  EC      (order_*.csv ※GIFT除く):  {ec_true:>6}")

    print("\n" + "=" * 55)
    print(f"[2] Excel の ORDER欄合計（実際値）")
    print("=" * 55)
    for c in channels:
        print(f"  {c:<12}: {actual[c]:>6}")

    print("\n" + "=" * 55)
    print("[3] 突き合わせ結果（正解値 vs 実際値）")
    print("=" * 55)

    all_ok = True

    checks = [
        ("ZOZO", zozo_true, actual["ZOZO"]),
        ("OIOI", oioi_true, actual["OIOI"]),
        ("EC", ec_true, actual["EC"])
    ] + [(s, store_totals[s], actual[s]) for s in stores_list]

    for name, truth, got in checks:
        ok = (truth == got)
        mark = "[OK]" if ok else "[NG] 合計不一致！"
        print(f"  {name:<12}: 正解={truth:>5}  実際={got:>5}  {mark}")
        if not ok:
            all_ok = False

    # 商品レベルのチェック結果の出力 (差異分析レポートの表示)
    if diff_details:
        all_ok = False
        print("\n" + "!" * 55)
        print(f"[NG] 商品レベル（セル単位）の不一致が {len(diff_details)} 件検出されました！ (差異分析レポート)")
        print("!" * 55)
        print(f"{'商品コード':<15} | {'チャネル/店舗':<10} | {'期待値(CSV)':<8} | {'実際値(Excel)':<9} | {'行番号':<5}")
        print("-" * 55)
        for diff in diff_details[:50]:
            print(f"{diff['code']:<15} | {diff['channel']:<10} | {diff['expected']:>8} | {diff['actual']:>9} | {diff['row']:>5}")
        if len(diff_details) > 50:
            print(f"... (他 {len(diff_details) - 50} 件の差異があります)")

    return all_ok
