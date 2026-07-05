"""
データ検証スクリプト: RETAIL_完成版.xlsx vs 元CSV の合計突き合わせ
"""
import pandas as pd
import openpyxl
import glob, os

result_path = r"C:\Users\kxnxg\OneDrive\デスクトップ\automation\RETAIL_完成版.xlsx"
input_dir = r"C:\Users\kxnxg\OneDrive\デスクトップ\automation\input"

def find_csv(keywords, exclude=None):
    for p in glob.glob(os.path.join(input_dir, "*.csv")):
        name = os.path.basename(p)
        if all(k in name for k in keywords):
            if exclude and any(e in name for e in exclude):
                continue
            return p
    return None

# ────────────────────────────────────────
# 1. 元CSVから「正解の合計」を計算
# ────────────────────────────────────────
print("=" * 55)
print("[1] 元CSV の合計（正解値）")
print("=" * 55)

# ZOZO
df_zozo = pd.read_csv(find_csv(["卸売上明細"], exclude=["卸売上明細 (1)", "卸売上明細(1)"]), encoding="cp932", dtype=str)
df_zozo = df_zozo[df_zozo["ブランド名"].isin(["FRV", "FJALLRAVEN"])].copy()
df_zozo["販売数量"] = pd.to_numeric(df_zozo["販売数量"], errors="coerce").fillna(0)
zozo_true = int(df_zozo["販売数量"].sum())
print(f"  ZOZO    (卸売上明細.csv):          {zozo_true:>6}")

# OIOI
df_oioi = pd.read_csv(find_csv(["卸売上明細 (1)"]), encoding="cp932", dtype=str)
df_oioi = df_oioi[df_oioi["ブランド名"].isin(["FRV", "FJALLRAVEN"])].copy()
df_oioi["販売数量"] = pd.to_numeric(df_oioi["販売数量"], errors="coerce").fillna(0)
oioi_true = int(df_oioi["販売数量"].sum())
print(f"  OIOI    (卸売上明細 (1).csv):      {oioi_true:>6}")

# 店舗
df_store = pd.read_csv(find_csv(["営業日付別売上分析"]), encoding="cp932", dtype=str)
df_store = df_store[df_store["表記部門名1"].isin(["FRV", "FJALLRAVEN"])].copy()
df_store["数量"] = pd.to_numeric(df_store["数量"], errors="coerce").fillna(0)
store_pivot = df_store.pivot_table(index="3rd Item No.", columns="店舗名称", values="数量", aggfunc="sum")

def store_total(keyword):
    cols = [c for c in store_pivot.columns if keyword in c or (keyword == "ヒュッテ" and "HUTTE" in c.upper())]
    return int(store_pivot[cols].sum().sum()) if cols else 0

store_totals = {
    "名古屋":     store_total("名古屋"),
    "TOKYO":     store_total("TOKYO"),
    "ルクア大阪": store_total("ルクア大阪"),
    "ヒュッテ":   store_total("ヒュッテ"),
    "京王新宿":   store_total("京王新宿"),
    "大丸心斎橋": store_total("大丸心斎橋"),
    "6142":      store_total("6142"),
    "玉川高島屋": store_total("玉川高島屋"),
    "NODE":      store_total("NODE"),
    "NARITA":    store_total("NARITA"),
}
for name, total in store_totals.items():
    print(f"  {name:<12}(営業日付別売上分析.csv): {total:>6}")

# EC
df_ec = pd.read_csv(find_csv(["order_"]), encoding="cp932", dtype=str)
df_ec = df_ec[~df_ec["決済方法(ステータス)"].str.contains("キャンセル", na=False)]
df_ec = df_ec[~df_ec["注文者"].str.contains("店舗客注", na=False)]
df_ec["個数"] = pd.to_numeric(df_ec["個数"], errors="coerce").fillna(0)
key_col = "オプション独自コード" if "オプション独自コード" in df_ec.columns else "商品コード"
# GIFT500はギフトラッピング料のためorderシート対象外（以前確認済み）
df_ec = df_ec[~df_ec[key_col].fillna("").str.contains("GIFT", na=False)]
ec_true = int(df_ec["個数"].sum())
print(f"  EC      (order_*.csv ※GIFT除く):  {ec_true:>6}")

# ────────────────────────────────────────
# 2. RETAIL_完成版.xlsx から実際の集計値
# ────────────────────────────────────────
print("\n" + "=" * 55)
print("[2] RETAIL_完成版.xlsx の ORDER欄合計（実際値）")
print("=" * 55)

wb = openpyxl.load_workbook(result_path, data_only=True)
ws = wb["order"]

def col_total(col_num):
    return sum(
        (ws.cell(row=r, column=col_num).value or 0)
        for r in range(4, ws.max_row + 1)
        if ws.cell(row=r, column=2).value is not None
    )

# 列番号: AE(31)=ZOZO, AF(32)=名古屋, AG(33)=OIOI, AH(34)=EC
# AI(35)=TOKYO, AJ(36)=ルクア大阪, AK(37)=ヒュッテ, AL(38)=京王新宿
# AM(39)=大丸心斎橋, AN(40)=6142, AO(41)=玉川高島屋, AP(42)=NODE, AQ(43)=NARITA
actual = {
    "ZOZO":      col_total(31),
    "名古屋":    col_total(32),
    "OIOI":      col_total(33),
    "EC":        col_total(34),
    "TOKYO":     col_total(35),
    "ルクア大阪": col_total(36),
    "ヒュッテ":  col_total(37),
    "京王新宿":  col_total(38),
    "大丸心斎橋": col_total(39),
    "6142":      col_total(40),
    "玉川高島屋": col_total(41),
    "NODE":      col_total(42),
    "NARITA":    col_total(43),
}
for name, val in actual.items():
    print(f"  {name:<12}: {val:>6}")

wb.close()

# ────────────────────────────────────────
# 3. 突き合わせ（差異チェック）
# ────────────────────────────────────────
print("\n" + "=" * 55)
print("[3] 突き合わせ結果（正解値 vs 実際値）")
print("=" * 55)

checks = [
    ("ZOZO",      zozo_true,             actual["ZOZO"]),
    ("OIOI",      oioi_true,             actual["OIOI"]),
    ("EC",        ec_true,               actual["EC"]),
    ("名古屋",    store_totals["名古屋"], actual["名古屋"]),
    ("TOKYO",     store_totals["TOKYO"],  actual["TOKYO"]),
    ("ルクア大阪",store_totals["ルクア大阪"], actual["ルクア大阪"]),
    ("ヒュッテ",  store_totals["ヒュッテ"],  actual["ヒュッテ"]),
    ("京王新宿",  store_totals["京王新宿"],  actual["京王新宿"]),
    ("大丸心斎橋",store_totals["大丸心斎橋"],actual["大丸心斎橋"]),
    ("6142",      store_totals["6142"],   actual["6142"]),
    ("玉川高島屋",store_totals["玉川高島屋"],actual["玉川高島屋"]),
    ("NODE",      store_totals["NODE"],   actual["NODE"]),
    ("NARITA",    store_totals["NARITA"], actual["NARITA"]),
]

all_ok = True
for name, truth, got in checks:
    ok = (truth == got)
    mark = "[OK]" if ok else "[NG] 不一致！"
    print(f"  {name:<12}: 正解={truth:>5}  実際={got:>5}  {mark}")
    if not ok:
        all_ok = False

print("\n" + ("[OK] すべて一致しています！" if all_ok else "[NG] 不一致があります。上記を確認してください。"))
